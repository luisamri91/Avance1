"""Genera un Word aparte con solo las historias de usuario (historias_usuario.json).

Página horizontal, márgenes de 2,54 cm y Times New Roman, listo para copiar la tabla
a otro documento. Uso (desde docs/):  python3 construir_historias.py
Para una versión nueva, subir VERSION en uno.
"""
import json

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from construir_documento import ALTO_CARTA, ANCHO_CARTA, HISTORIAS, MARGEN, construir_tabla_historias

VERSION = 1
SALIDA = f"Avance1_Historias_de_Usuario_v{VERSION}.docx"


def main():
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.size = Pt(12)
    propiedades = normal.element.get_or_add_rPr()
    fuentes = propiedades.find(qn("w:rFonts"))
    if fuentes is None:
        fuentes = OxmlElement("w:rFonts")
        propiedades.insert(0, fuentes)
    for lado in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fuentes.set(qn(lado), "Times New Roman")
    seccion = doc.sections[0]
    seccion.orientation = WD_ORIENT.LANDSCAPE
    seccion.page_width, seccion.page_height = ALTO_CARTA, ANCHO_CARTA
    seccion.left_margin = seccion.right_margin = MARGEN
    seccion.top_margin = seccion.bottom_margin = MARGEN

    titulo = doc.add_paragraph()
    trozo = titulo.add_run("Historias de usuario")
    trozo.bold = True
    titulo.paragraph_format.space_after = Pt(12)

    with open(HISTORIAS, encoding="utf-8") as archivo:
        historias = json.load(archivo)
    tabla = construir_tabla_historias(doc, historias)
    tabla._tbl.addnext(OxmlElement("w:p"))  # Word necesita un párrafo después de la tabla
    doc.save(SALIDA)
    print("Generado", SALIDA, "con", len(historias), "historias")


if __name__ == "__main__":
    main()
