"""Arma el Word con las capturas del sistema funcionando (docs/capturas_sistema/*.png)."""
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from PIL import Image

VERSION = 1
CARPETA = Path(__file__).parent / "capturas_sistema"

# archivo (sin número) -> (sección, pie de figura)
PIES = {
    "servidor_iniciado": ("Servidor", "HU 12. Ventana del servidor recién iniciada en el puerto 5000 con la base Derby."),
    "login_error": ("Inicio de sesión", "HU 1. Contraseña incorrecta: el sistema no dice qué dato falló y limpia la contraseña."),
    "login": ("Inicio de sesión", "HU 1. Pantalla de inicio de sesión de una caja."),
    "menu_administrador": ("Menú principal", "Menú del administrador con el resumen del día y el aviso de productos bajo el stock mínimo (HU 7)."),
    "clientes_formulario": ("Clientes", "HU 3. Registro de un cliente nuevo con cédula jurídica."),
    "clientes_creado": ("Clientes", "HU 3. Mensaje del servidor al guardar el cliente."),
    "clientes_correo_invalido": ("Clientes", "HU 3. Datos inválidos: el servidor rechaza el correo y conserva los demás datos."),
    "clientes_eliminar": ("Clientes", "HU 3. Eliminar un cliente sin facturas: confirmación y resultado."),
    "clientes_lista": ("Clientes", "HU 3. Lista de clientes después de crear, editar el teléfono y eliminar."),
    "productos_bebida": ("Productos e inventario", "HU 5. Producto nuevo de tipo Bebida: el recuadro cambia y avisa el IVA de 13 %."),
    "productos_creado": ("Productos e inventario", "HU 5. Producto registrado."),
    "ajuste_stock": ("Productos e inventario", "HU 6. Ajuste de stock: compra de 50 sacos de arroz (18 a 68 unidades)."),
    "ajuste_aplicado": ("Productos e inventario", "HU 6. Confirmación y resultado del ajuste."),
    "productos_lista": ("Productos e inventario", "HU 5 y 7. Catálogo ordenado por nombre (Comparable) con el estado de cada producto."),
    "categorias": ("Productos e inventario", "HU 4. Pestaña Categorías con la categoría nueva Snacks."),
    "bajo_el_minimo": ("Productos e inventario", "HU 7. Pestaña Bajo el mínimo, de menos a más existencias (Comparator)."),
    "menu_vendedor": ("Menú principal", "HU 2. Menú de un VENDEDOR: no aparece la opción Usuarios."),
    "factura_stock_insuficiente": ("Facturación", "HU 8. Pedir más unidades de las que hay: la línea no se agrega."),
    "nueva_factura": ("Facturación", "HU 8. Factura con tres productos, la cola agregada dos veces en la misma línea, descuento de 2 % y transferencia."),
    "vista_factura": ("Facturación", "HU 8 y 10. Vista previa de la factura FC-000001 emitida por el servidor."),
    "factura_archivo_guardado": ("Facturación", "HU 10. Archivo FC-000001.txt guardado."),
    "caja1_stock_insuficiente": ("Varias cajas al mismo tiempo", "HU 9. Las dos cajas querían las últimas 9 unidades de detergente; la caja de Ana emitió primero y la de Carlos recibe este mensaje."),
    "caja1_conserva_factura": ("Varias cajas al mismo tiempo", "HU 9. La caja que perdió conserva su factura para corregirla."),
    "historial_anular": ("Historial", "HU 11. Anulación de la factura FC-000002: confirmación y resultado."),
    "historial": ("Historial", "HU 11. Historial con la factura anulada en rojo y el total del periodo solo con las emitidas."),
    "usuarios_formulario": ("Usuarios", "HU 2. Creación de un usuario vendedor."),
    "usuarios_creado": ("Usuarios", "HU 2. Usuario creado."),
    "usuarios_desactivar": ("Usuarios", "HU 2. Eliminar a un usuario con facturas: se desactiva para conservar el historial."),
    "usuarios_lista": ("Usuarios", "HU 2. Lista de usuarios: cjimenez queda Inactivo y lmora aparece nuevo."),
    "servidor_bitacora": ("Servidor", "HU 12. Tres cajas conectadas, cada una en su hilo, y la bitácora de operaciones."),
    "bd_facturas": ("Base de datos", "Tabla FACTURAS leída directamente de Derby: la FC-000002 quedó ANULADA."),
    "bd_productos": ("Base de datos", "Tabla PRODUCTOS con las existencias actualizadas por ajustes, ventas y la anulación."),
    "bd_clientes": ("Base de datos", "Tabla CLIENTES con el cliente nuevo y el teléfono editado."),
    "escritorio_servidor_y_tres_cajas": ("Servidor", "El servidor y tres cajas abiertas al mismo tiempo (administradora y dos vendedores)."),
    "conexion_perdida": ("Servidor", "HU 12. Al detener el servidor, las cajas muestran «Se perdió la conexión con el servidor.»"),
}

doc = Document()
sec = doc.sections[0]
sec.left_margin = sec.right_margin = Cm(2)
sec.top_margin = sec.bottom_margin = Cm(2)
estilo = doc.styles["Normal"]
estilo.font.name = "Calibri"
estilo.font.size = Pt(11)

doc.add_heading(f"Fidecompro: sistema funcionando (capturas v{VERSION})", 0)
doc.add_paragraph(
    "Capturas de las ventanas reales del código v6 corriendo: el servidor con su pool de hilos y la base de datos "
    "Apache Derby, y tres cajas cliente (una administradora y dos vendedores) conectadas al mismo tiempo por sockets. "
    "Se recorren las 12 historias de usuario: CRUD de clientes, productos, categorías y usuarios, ajustes de stock, "
    "facturación con varias cajas, historial, anulación y la ventana del servidor.")
doc.add_paragraph(
    "Prueba de concurrencia aparte (HU 9): dos conexiones en dos hilos pidieron al mismo tiempo las 50 botellas de agua "
    "que había; el servidor emitió una sola factura, la otra recibió «Stock insuficiente de Agua 600 ml (caja 24): hay 0 "
    "unidades.» y las existencias quedaron en 0, nunca en negativo.")

seccion_actual = None
figura = 0
for png in sorted(CARPETA.glob("*.png")):
    clave = png.stem.split("_", 1)[1]
    seccion, pie = PIES.get(clave, ("Otras", clave))
    if seccion != seccion_actual:
        doc.add_heading(seccion, 1)
        seccion_actual = seccion
    figura += 1
    ancho = 17 if "escritorio" in clave or "conexion" in clave else 15
    w, h = Image.open(png).size
    if w < 600:
        ancho = 8
    doc.add_picture(str(png), width=Cm(ancho))
    imagen = doc.paragraphs[-1]
    imagen.alignment = WD_ALIGN_PARAGRAPH.CENTER
    imagen.paragraph_format.keep_with_next = True
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(14)
    p.add_run(f"Figura {figura}. ").bold = True
    p.add_run(pie).italic = True

salida = Path(__file__).parent / f"Fidecompro_Sistema_Funcionando_v{VERSION}.docx"
doc.save(salida)
print(salida)
