"""Genera el documento de diseño en Word y PDF a partir de Avance1_Documento_Diseno.md.

Uso (desde la carpeta docs):  python3 construir_documento.py

1. Convierte el Markdown a .docx con pandoc.
2. Da formato: tamaño carta, imágenes al ancho de la página y tablas con bordes.
3. Reemplaza los marcadores TABLA_REQUERIMIENTOS_… por las tablas de
   requerimientos (requerimientos.json) y TABLA_HISTORIAS_USUARIO por la tabla de
   historias de usuario (historias_usuario.json), esta en una sección horizontal.
4. Convierte el .docx a PDF con LibreOffice.

El nombre de los archivos lleva la versión que dice el subtítulo del Markdown
("Versión N"), por ejemplo Avance1_Documento_Diseno_v4.docx.
"""
import copy
import json
import re
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

CARPETA = Path(__file__).resolve().parent
FUENTE = CARPETA / "Avance1_Documento_Diseno.md"
HISTORIAS = CARPETA / "historias_usuario.json"
REQUERIMIENTOS = CARPETA / "requerimientos.json"
MARCADOR = "TABLA_HISTORIAS_USUARIO"
MARCADORES_REQUERIMIENTOS = {
    "TABLA_REQUERIMIENTOS_FUNCIONALES": "funcionales",
    "TABLA_REQUERIMIENTOS_NO_FUNCIONALES": "no_funcionales",
}

ANCHO_CARTA, ALTO_CARTA = Cm(21.59), Cm(27.94)
AZUL_ENCABEZADO = "1F3864"
GRIS_SEPARADOR = "D9D9D9"
AZUL_REQUERIMIENTOS = "1F4E79"
CELESTE_FILA = "BDD7EE"
COLUMNAS_REQUERIMIENTOS = [("#", 0.9), ("Requerimiento", 5.0), ("Descripción", 9.4), ("Prioridad", 2.2)]
COLUMNAS_HU = [
    ("Identificador (ID) de la historia", 1.7),
    ("Rol", 2.1),
    ("Característica / Funcionalidad", 2.9),
    ("Razón / Resultado", 2.9),
    ("Número (#) de escenario", 1.3),
    ("Criterio de aceptación (Título)", 2.6),
    ("Contexto", 4.0),
    ("Evento", 3.0),
    ("Resultado / Comportamiento esperado", 4.4),
]


def leer_version():
    texto = FUENTE.read_text(encoding="utf-8")
    encontrado = re.search(r"Versión (\d+)", texto)
    if not encontrado:
        sys.exit("No se encontró 'Versión N' en el subtítulo del documento")
    return encontrado.group(1)


def sombrear(celda, color):
    sombra = OxmlElement("w:shd")
    sombra.set(qn("w:val"), "clear")
    sombra.set(qn("w:fill"), color)
    celda._tc.get_or_add_tcPr().append(sombra)


def poner_bordes(tabla, color="8E9BAB"):
    bordes = OxmlElement("w:tblBorders")
    for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
        borde = OxmlElement(f"w:{lado}")
        borde.set(qn("w:val"), "single")
        borde.set(qn("w:sz"), "4")
        borde.set(qn("w:color"), color)
        bordes.append(borde)
    tabla._tbl.tblPr.append(bordes)


def dar_formato_general(doc):
    seccion = doc.sections[0]
    seccion.page_width, seccion.page_height = ANCHO_CARTA, ALTO_CARTA
    seccion.left_margin = seccion.right_margin = Cm(2)
    ancho_max, alto_max = Cm(17.5), Cm(20.5)  # deja lugar al título y al pie de figura
    for imagen in doc.inline_shapes:
        proporcion = imagen.height / imagen.width
        ancho, alto = ancho_max, int(ancho_max * proporcion)
        if alto > alto_max:
            alto, ancho = alto_max, int(alto_max / proporcion)
        imagen.width, imagen.height = int(ancho), int(alto)
    for tabla in doc.tables:
        poner_bordes(tabla)
        for celda in tabla.rows[0].cells:
            sombrear(celda, "DCE4EE")
            for parrafo in celda.paragraphs:
                for trozo in parrafo.runs:
                    trozo.bold = True
        for fila in tabla.rows:
            for celda in fila.cells:
                for parrafo in celda.paragraphs:
                    for trozo in parrafo.runs:
                        trozo.font.size = Pt(9.5)


def escribir(celda, texto, negrita=False, color=None, centrado=False, tamano=8.5):
    parrafo = celda.paragraphs[0]
    parrafo.paragraph_format.space_after = Pt(0)
    if centrado:
        parrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    trozo = parrafo.add_run(texto)
    trozo.font.size = Pt(tamano)
    trozo.bold = negrita
    if color:
        trozo.font.color.rgb = RGBColor.from_string(color)


def construir_tabla_historias(doc, historias):
    filas = 1 + sum(len(h["escenarios"]) + 1 for h in historias)
    tabla = doc.add_table(rows=filas, cols=len(COLUMNAS_HU))
    tabla.alignment = WD_TABLE_ALIGNMENT.CENTER
    tabla.autofit = False
    poner_bordes(tabla, "000000")

    encabezado = tabla.rows[0]
    encabezado_pr = encabezado._tr.get_or_add_trPr()
    repetir = OxmlElement("w:tblHeader")  # repite el encabezado en cada página
    encabezado_pr.append(repetir)
    for i, (titulo, _) in enumerate(COLUMNAS_HU):
        celda = encabezado.cells[i]
        sombrear(celda, AZUL_ENCABEZADO)
        celda.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        escribir(celda, titulo, color="FFFFFF", centrado=True)

    fila = 1
    for historia in historias:
        inicio = fila
        for numero, (titulo, contexto, evento, resultado) in enumerate(historia["escenarios"], start=1):
            celdas = tabla.rows[fila].cells
            escribir(celdas[4], str(numero), centrado=True)
            escribir(celdas[5], titulo)
            escribir(celdas[6], contexto)
            escribir(celdas[7], evento)
            escribir(celdas[8], resultado)
            fila += 1
        # Las cuatro primeras columnas se combinan en todos los escenarios de la historia
        for columna, clave in enumerate(("id", "rol", "caracteristica", "razon")):
            combinada = tabla.cell(inicio, columna).merge(tabla.cell(fila - 1, columna))
            for sobrante in combinada.paragraphs[1:]:  # merge deja un párrafo vacío por celda
                sobrante._p.getparent().remove(sobrante._p)
            escribir(combinada, historia[clave])
        for celda in tabla.rows[fila].cells:  # fila gris que separa una historia de otra
            sombrear(celda, GRIS_SEPARADOR)
        fila += 1

    fijar_anchos(tabla, COLUMNAS_HU)
    return tabla


def construir_tabla_requerimientos(doc, requerimientos):
    tabla = doc.add_table(rows=1 + len(requerimientos), cols=len(COLUMNAS_REQUERIMIENTOS))
    tabla.alignment = WD_TABLE_ALIGNMENT.CENTER
    tabla.autofit = False
    poner_bordes(tabla, "000000")
    encabezado = tabla.rows[0]
    encabezado._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for i, (titulo, _) in enumerate(COLUMNAS_REQUERIMIENTOS):
        sombrear(encabezado.cells[i], AZUL_REQUERIMIENTOS)
        escribir(encabezado.cells[i], titulo, negrita=True, color="FFFFFF", tamano=10)
    for numero, (requerimiento, descripcion, prioridad) in enumerate(requerimientos, start=1):
        celdas = tabla.rows[numero].cells
        for celda, texto, negrita in zip(celdas, (str(numero), requerimiento, descripcion, prioridad),
                                         (True, False, False, False)):
            escribir(celda, texto, negrita=negrita, tamano=10)
            if numero % 2 == 1:  # filas alternas en celeste
                sombrear(celda, CELESTE_FILA)
    fijar_anchos(tabla, COLUMNAS_REQUERIMIENTOS)
    return tabla


def insertar_tablas_requerimientos(doc):
    with open(REQUERIMIENTOS, encoding="utf-8") as archivo:
        requerimientos = json.load(archivo)
    for marcador_texto, clave in MARCADORES_REQUERIMIENTOS.items():
        marcador = next(p for p in doc.paragraphs if p.text.strip() == marcador_texto)
        tabla = construir_tabla_requerimientos(doc, requerimientos[clave])
        marcador._p.addnext(tabla._tbl)
        marcador._p.getparent().remove(marcador._p)


def fijar_anchos(tabla, columnas):
    # Word y LibreOffice respetan los anchos solo si la tabla tiene diseño fijo y la cuadrícula los declara
    propiedades = tabla._tbl.tblPr
    diseno = OxmlElement("w:tblLayout")
    diseno.set(qn("w:type"), "fixed")
    propiedades.append(diseno)
    ancho_total = OxmlElement("w:tblW")
    ancho_total.set(qn("w:w"), str(sum(Cm(a).twips for _, a in columnas)))
    ancho_total.set(qn("w:type"), "dxa")
    for viejo in propiedades.findall(qn("w:tblW")):
        propiedades.remove(viejo)
    propiedades.append(ancho_total)
    for columna, (_, ancho) in zip(tabla._tbl.tblGrid.findall(qn("w:gridCol")), columnas):
        columna.set(qn("w:w"), str(Cm(ancho).twips))
    for fila_tabla in tabla.rows:
        for i, (_, ancho) in enumerate(columnas):
            fila_tabla.cells[i].width = Cm(ancho)


def propiedades_de_seccion(base, horizontal):
    seccion = copy.deepcopy(base)
    tamano = seccion.find(qn("w:pgSz"))
    margenes = seccion.find(qn("w:pgMar"))
    if horizontal:
        tamano.set(qn("w:w"), str(ALTO_CARTA.twips))
        tamano.set(qn("w:h"), str(ANCHO_CARTA.twips))
        tamano.set(qn("w:orient"), "landscape")
        for lado in ("w:left", "w:right"):
            margenes.set(qn(lado), str(Cm(1.5).twips))
        for lado in ("w:top", "w:bottom"):
            margenes.set(qn(lado), str(Cm(1.8).twips))
    return seccion


def insertar_tabla_historias(doc):
    with open(HISTORIAS, encoding="utf-8") as archivo:
        historias = json.load(archivo)
    marcador = next(p for p in doc.paragraphs if p.text.strip() == MARCADOR)
    titulo = next(p for p in doc.paragraphs if re.match(r"\d+\. Historias de usuario", p.text.strip()))
    seccion_base = doc.sections[0]._sectPr

    # Un párrafo vacío antes del título de la sección 3 cierra la parte vertical,
    # así el título, la introducción y la tabla quedan juntos en páginas horizontales
    corte = OxmlElement("w:p")
    corte_pr = OxmlElement("w:pPr")
    corte_pr.append(propiedades_de_seccion(seccion_base, horizontal=False))
    corte.append(corte_pr)
    titulo._p.addprevious(corte)

    tabla = construir_tabla_historias(doc, historias)
    marcador._p.addnext(tabla._tbl)
    marcador._p.getparent().remove(marcador._p)

    # Un párrafo vacío después de la tabla cierra la sección horizontal
    cierre = OxmlElement("w:p")
    cierre_pr = OxmlElement("w:pPr")
    cierre_pr.append(propiedades_de_seccion(seccion_base, horizontal=True))
    cierre.append(cierre_pr)
    tabla._tbl.addnext(cierre)


def main():
    version = leer_version()
    salida = CARPETA / f"Avance1_Documento_Diseno_v{version}.docx"
    subprocess.run(["pandoc", FUENTE.name, "-o", salida.name, "--resource-path=."],
                   cwd=CARPETA, check=True)
    doc = Document(salida)
    dar_formato_general(doc)
    insertar_tablas_requerimientos(doc)
    insertar_tabla_historias(doc)
    doc.save(salida)
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", salida.name],
                   cwd=CARPETA, check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    print(f"Generados {salida.name} y {salida.with_suffix('.pdf').name}")


if __name__ == "__main__":
    main()
