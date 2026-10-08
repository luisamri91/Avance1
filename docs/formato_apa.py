"""Formato APA 7 para los Word que generan construir_documento.py y construir_diagrama_clases.py.

* Times New Roman 12, texto con interlineado doble y sangría de 1,27 cm.
* Títulos de nivel 1 (centrado y negrita), 2 (izquierda y negrita) y 3 (izquierda, negrita y cursiva).
* Portada de estudiante: título en negrita, centrado, y los datos debajo; el resto empieza en otra página.
* Número de página arriba a la derecha en todas las páginas.
* Tablas: "Tabla N" en negrita y el título en cursiva arriba; solo líneas horizontales.
* Figuras: "Figura N" en negrita y el título en cursiva arriba de la imagen.
"""
import copy

from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

LETRA = "Times New Roman"
TAMANO = Pt(12)
SANGRIA = Cm(1.27)
NEGRO = RGBColor(0, 0, 0)


def estilo(doc, nombre):
    """Busca un estilo por nombre sin importar mayúsculas ("heading 1" o "Heading 1")."""
    for candidato in doc.styles:
        if candidato.name and candidato.name.lower() == nombre.lower():
            return candidato
    return None


def poner_letra(estilo_docx):
    rpr = estilo_docx.element.get_or_add_rPr()
    fuentes = rpr.find(qn("w:rFonts"))
    if fuentes is None:
        fuentes = OxmlElement("w:rFonts")
        rpr.insert(0, fuentes)
    for atributo in list(fuentes.attrib):
        if "Theme" in atributo:  # las fuentes del tema le ganan a la letra elegida
            del fuentes.attrib[atributo]
    for lado in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fuentes.set(qn(lado), LETRA)


def interlineado(formato, doble):
    formato.line_spacing_rule = WD_LINE_SPACING.DOUBLE if doble else WD_LINE_SPACING.SINGLE
    formato.space_before = Pt(0)
    formato.space_after = Pt(0)


def dar_estilos(doc):
    for estilo_docx in doc.styles:
        if estilo_docx.type == 1:  # estilos de párrafo
            poner_letra(estilo_docx)
    texto = estilo(doc, "Body Text")
    for nombre in ("Normal", "Body Text", "First Paragraph"):
        e = estilo(doc, nombre)
        if e is not None:
            e.font.size = TAMANO
            e.font.color.rgb = NEGRO
    for nombre in ("Body Text", "First Paragraph"):
        e = estilo(doc, nombre)
        if e is not None:
            interlineado(e.paragraph_format, doble=True)
            e.paragraph_format.first_line_indent = SANGRIA
    compacto = estilo(doc, "Compact")  # celdas de tablas y listas: sencillo y sin sangría
    if compacto is not None:
        interlineado(compacto.paragraph_format, doble=False)
        compacto.paragraph_format.first_line_indent = Cm(0)
    niveles = {"Heading 1": (WD_ALIGN_PARAGRAPH.CENTER, False),
               "Heading 2": (WD_ALIGN_PARAGRAPH.LEFT, False),
               "Heading 3": (WD_ALIGN_PARAGRAPH.LEFT, True)}
    for nombre, (alineacion, cursiva) in niveles.items():
        e = estilo(doc, nombre)
        if e is None:
            continue
        e.font.size = TAMANO
        e.font.bold = True
        e.font.italic = cursiva
        e.font.color.rgb = NEGRO
        e.paragraph_format.alignment = alineacion
        e.paragraph_format.keep_with_next = True
        interlineado(e.paragraph_format, doble=True)
        e.paragraph_format.space_before = Pt(12)  # separa el título de una tabla que termine justo antes
    for nombre in ("Title", "Subtitle", "Author", "Date"):
        e = estilo(doc, nombre)
        if e is None:
            continue
        e.font.size = TAMANO
        e.font.color.rgb = NEGRO
        e.font.bold = nombre in ("Title", "Subtitle")
        e.font.italic = False
        e.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        interlineado(e.paragraph_format, doble=True)
    for nombre in ("Table Caption", "Image Caption"):
        e = estilo(doc, nombre)
        if e is not None:
            e.font.size = TAMANO
            e.font.color.rgb = NEGRO
            e.font.bold = False
            e.font.italic = True
            e.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
            e.paragraph_format.keep_with_next = True
            interlineado(e.paragraph_format, doble=True)
    if texto is None:
        raise SystemExit("El .docx no tiene el estilo Body Text de pandoc")


def armar_portada(doc):
    """Baja el título como pide APA y hace que el documento empiece en la página 2."""
    titulo = next(p for p in doc.paragraphs if p.style.name == "Title")
    titulo.paragraph_format.space_before = Pt(96)
    fecha = next(p for p in doc.paragraphs if p.style.name == "Date")
    fecha.add_run().add_break(WD_BREAK.PAGE)


def poner_numero_pagina(seccion):
    encabezado = seccion.header
    encabezado.is_linked_to_previous = False
    parrafo = encabezado.paragraphs[0]
    parrafo.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    campo = OxmlElement("w:fldSimple")
    campo.set(qn("w:instr"), "PAGE")
    trozo = OxmlElement("w:r")
    propiedades = OxmlElement("w:rPr")
    fuentes = OxmlElement("w:rFonts")
    for lado in ("w:ascii", "w:hAnsi"):
        fuentes.set(qn(lado), LETRA)
    tamano = OxmlElement("w:sz")
    tamano.set(qn("w:val"), "24")
    propiedades.extend([fuentes, tamano])
    texto = OxmlElement("w:t")
    texto.text = "1"
    trozo.extend([propiedades, texto])
    campo.append(trozo)
    parrafo._p.append(campo)


def bordes_apa(tabla):
    """Solo líneas horizontales: arriba, debajo de los encabezados, abajo y una línea fina entre filas."""
    propiedades = tabla._tbl.tblPr
    for viejo in propiedades.findall(qn("w:tblBorders")):
        propiedades.remove(viejo)
    bordes = OxmlElement("w:tblBorders")
    for lado, valor, grosor, color in (("top", "single", "8", "000000"), ("left", "nil", "0", "auto"),
                                       ("bottom", "single", "8", "000000"), ("right", "nil", "0", "auto"),
                                       ("insideH", "single", "2", "BFBFBF"), ("insideV", "nil", "0", "auto")):
        borde = OxmlElement(f"w:{lado}")
        borde.set(qn("w:val"), valor)
        borde.set(qn("w:sz"), grosor)
        borde.set(qn("w:color"), color)
        bordes.append(borde)
    propiedades.append(bordes)
    for celda in tabla.rows[0].cells:
        tc_pr = celda._tc.get_or_add_tcPr()
        bordes_celda = OxmlElement("w:tcBorders")
        abajo = OxmlElement("w:bottom")
        abajo.set(qn("w:val"), "single")
        abajo.set(qn("w:sz"), "8")
        abajo.set(qn("w:color"), "000000")
        bordes_celda.append(abajo)
        tc_pr.append(bordes_celda)


def etiqueta_antes(parrafo, texto):
    """Inserta antes de `parrafo` uno nuevo, del mismo estilo, con `texto` en negrita y sin cursiva."""
    nuevo = copy.deepcopy(parrafo._p)
    for hijo in list(nuevo):
        if hijo.tag != qn("w:pPr"):
            nuevo.remove(hijo)
    trozo = OxmlElement("w:r")
    propiedades = OxmlElement("w:rPr")
    for etiqueta, valor in (("w:b", None), ("w:i", "0")):
        elemento = OxmlElement(etiqueta)
        if valor:
            elemento.set(qn("w:val"), valor)
        propiedades.append(elemento)
    texto_xml = OxmlElement("w:t")
    texto_xml.text = texto
    trozo.extend([propiedades, texto_xml])
    nuevo.append(trozo)
    parrafo._p.addprevious(nuevo)


def numerar_tablas_y_figuras(doc):
    tablas = figuras = 0
    for parrafo in list(doc.paragraphs):
        nombre = parrafo.style.name
        if nombre == "Table Caption":
            tablas += 1
            etiqueta_antes(parrafo, f"Tabla {tablas}")
        elif nombre == "Image Caption":
            figuras += 1
            imagen = parrafo._p.getprevious()
            imagen.addprevious(parrafo._p)  # APA pone el título arriba de la figura
            etiqueta_antes(parrafo, f"Figura {figuras}")
    return tablas, figuras


def aplicar_apa(doc):
    dar_estilos(doc)
    armar_portada(doc)
    return numerar_tablas_y_figuras(doc)
